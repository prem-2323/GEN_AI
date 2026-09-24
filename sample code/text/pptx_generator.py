import io
from typing import Dict, Any, List
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.enum.text import PP_ALIGN
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE

# Theme Palettes
THEME_PALETTES = {
    "spotify_emerald": {
        "dark_bg": RGBColor(18, 18, 18),
        "title_dark": RGBColor(255, 255, 255),
        "subtitle_dark": RGBColor(179, 179, 179),
        "light_bg": RGBColor(24, 24, 24),
        "header": RGBColor(255, 255, 255),
        "accent": RGBColor(30, 215, 96),       # Spotify #1ED760
        "text_body": RGBColor(220, 220, 220),
        "muted": RGBColor(140, 140, 140),
        "box_bg": RGBColor(32, 32, 32),
        "card_border": RGBColor(50, 50, 50),
        "is_dark_theme": True,
    },
    "executive_navy": {
        "dark_bg": RGBColor(10, 25, 47),        # Navy 950
        "title_dark": RGBColor(255, 255, 255),
        "subtitle_dark": RGBColor(148, 163, 184),
        "light_bg": RGBColor(248, 250, 252),
        "header": RGBColor(15, 23, 42),
        "accent": RGBColor(14, 165, 233),       # Sky 500
        "text_body": RGBColor(51, 65, 85),
        "muted": RGBColor(100, 116, 139),
        "box_bg": RGBColor(255, 255, 255),
        "card_border": RGBColor(226, 232, 240),
        "is_dark_theme": False,
    },
    "corporate_purple": {
        "dark_bg": RGBColor(30, 27, 75),        # Indigo 950
        "title_dark": RGBColor(255, 255, 255),
        "subtitle_dark": RGBColor(199, 210, 254),
        "light_bg": RGBColor(248, 250, 252),
        "header": RGBColor(30, 41, 59),
        "accent": RGBColor(99, 102, 241),       # Indigo 500
        "text_body": RGBColor(51, 65, 85),
        "muted": RGBColor(100, 116, 139),
        "box_bg": RGBColor(255, 255, 255),
        "card_border": RGBColor(226, 232, 240),
        "is_dark_theme": False,
    },
    "crimson_minimal": {
        "dark_bg": RGBColor(24, 24, 27),        # Zinc 900
        "title_dark": RGBColor(255, 255, 255),
        "subtitle_dark": RGBColor(161, 161, 170),
        "light_bg": RGBColor(250, 250, 250),
        "header": RGBColor(24, 24, 27),
        "accent": RGBColor(244, 63, 94),        # Rose 500
        "text_body": RGBColor(63, 63, 70),
        "muted": RGBColor(113, 113, 122),
        "box_bg": RGBColor(255, 255, 255),
        "card_border": RGBColor(228, 228, 231),
        "is_dark_theme": False,
    },
    "modern_tech": {
        "dark_bg": RGBColor(15, 23, 42),        # Slate 900
        "title_dark": RGBColor(255, 255, 255),  # White
        "subtitle_dark": RGBColor(148, 163, 184),# Slate 400
        "light_bg": RGBColor(248, 250, 252),    # Slate 50
        "header": RGBColor(30, 41, 59),          # Slate 800
        "accent": RGBColor(37, 99, 235),         # Blue 600
        "text_body": RGBColor(51, 65, 85),       # Slate 700
        "muted": RGBColor(100, 116, 139),       # Slate 500
        "box_bg": RGBColor(255, 255, 255),      # White
        "card_border": RGBColor(226, 232, 240),  # Slate 200
        "is_dark_theme": False,
    }
}


def hex_to_rgb(hex_str: str, default_rgb=(30, 215, 96)) -> RGBColor:
    """Convert hex string (#1ED760 or 1ED760) into pptx RGBColor."""
    if not hex_str or not isinstance(hex_str, str):
        return RGBColor(*default_rgb)
    hex_clean = hex_str.strip().lstrip('#')
    if len(hex_clean) == 6:
        try:
            return RGBColor(int(hex_clean[0:2], 16), int(hex_clean[2:4], 16), int(hex_clean[4:6], 16))
        except ValueError:
            pass
    return RGBColor(*default_rgb)


def _resolve_palette(presentation_data: Dict[str, Any], theme_name: str = "spotify_emerald") -> dict:
    """Resolve theme palette with optional custom brand_voice primary/secondary colors."""
    brand_voice = presentation_data.get("brand_voice") or {}
    primary_hex = brand_voice.get("primary_color") or presentation_data.get("primary_color")
    secondary_hex = brand_voice.get("secondary_color") or presentation_data.get("secondary_color")

    if primary_hex and isinstance(primary_hex, str) and primary_hex.startswith("#"):
        accent_rgb = hex_to_rgb(primary_hex, (30, 215, 96))
        dark_rgb = hex_to_rgb(secondary_hex, (18, 18, 18)) if secondary_hex and secondary_hex.startswith("#") else RGBColor(18, 18, 18)
        box_rgb = RGBColor(min(255, dark_rgb[0] + 16), min(255, dark_rgb[1] + 16), min(255, dark_rgb[2] + 16))
        border_rgb = RGBColor(min(255, dark_rgb[0] + 36), min(255, dark_rgb[1] + 36), min(255, dark_rgb[2] + 36))
        return {
            "dark_bg": dark_rgb,
            "title_dark": RGBColor(255, 255, 255),
            "subtitle_dark": RGBColor(179, 179, 179),
            "light_bg": dark_rgb,
            "header": RGBColor(255, 255, 255),
            "accent": accent_rgb,
            "text_body": RGBColor(220, 220, 220),
            "muted": RGBColor(140, 140, 140),
            "box_bg": box_rgb,
            "card_border": border_rgb,
            "is_dark_theme": True,
        }

    chosen = theme_name or presentation_data.get("theme", "spotify_emerald")
    return THEME_PALETTES.get(chosen, THEME_PALETTES["spotify_emerald"])


def create_pptx_presentation(presentation_data: Dict[str, Any], theme_name: str = "spotify_emerald") -> io.BytesIO:
    """
    Generate a modern, executive 16:9 PowerPoint (.pptx) file from structured presentation JSON.
    Includes custom layouts, bullet points, two-column grids, visual recommendations, and speaker notes.
    """
    palette = _resolve_palette(presentation_data, theme_name)
    brand_voice = presentation_data.get("brand_voice") or {}
    disclaimer = brand_voice.get("disclaimer") or presentation_data.get("disclaimer", "")

    prs = Presentation()
    
    # Set 16:9 Widescreen dimensions (13.333 x 7.5 inches)
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    
    blank_layout = prs.slide_layouts[6]  # Blank slide layout
    
    title_text = presentation_data.get("presentation_title") or presentation_data.get("title") or "Executive Presentation"
    subtitle_text = presentation_data.get("subtitle") or presentation_data.get("main_message") or ""
    slides_data = presentation_data.get("slides") or []

    # If no slides array was returned, construct slide from raw contents
    if not slides_data:
        slides_data = [
            {
                "slide_number": 1,
                "title": title_text,
                "layout": "title",
                "subtitle": subtitle_text,
                "speaker_notes": "Welcome audience to the presentation."
            }
        ]

    # --- 1. TITLE SLIDE (Cover Slide) ---
    first_slide_data = slides_data[0]
    cover_slide = prs.slides.add_slide(blank_layout)
    _build_title_slide(cover_slide, title_text, subtitle_text, first_slide_data.get("speaker_notes", ""), palette)

    # --- 2. CONTENT SLIDES ---
    content_slides = slides_data[1:] if first_slide_data.get("layout") == "title" else slides_data
    
    for idx, slide_info in enumerate(content_slides, start=2):
        slide = prs.slides.add_slide(blank_layout)
        layout_type = str(slide_info.get("layout", "bullet_points")).lower()
        
        if "two_column" in layout_type or "grid" in layout_type or "comparison" in layout_type:
            _build_two_column_slide(slide, slide_info, palette, disclaimer=disclaimer)
        elif "quote" in layout_type or "callout" in layout_type:
            _build_quote_slide(slide, slide_info, palette, disclaimer=disclaimer)
        elif "metric" in layout_type or "kpi" in layout_type:
            _build_metrics_slide(slide, slide_info, palette, disclaimer=disclaimer)
        else:
            _build_bullet_slide(slide, slide_info, palette, disclaimer=disclaimer)
            
        # Attach Speaker Notes
        notes = slide_info.get("speaker_notes", "")
        if notes and hasattr(slide, "notes_slide"):
            slide.notes_slide.notes_text_frame.text = notes

    output_stream = io.BytesIO()
    prs.save(output_stream)
    output_stream.seek(0)
    return output_stream


def _build_title_slide(slide, title: str, subtitle: str, speaker_notes: str, palette: dict):
    """Build a striking Cover Title Slide."""
    bg = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0), Inches(0), Inches(13.333), Inches(7.5))
    bg.fill.solid()
    bg.fill.fore_color.rgb = palette["dark_bg"]
    bg.line.fill.background()

    # Accent Header Bar
    bar = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(1), Inches(2.2), Inches(0.15), Inches(3.2))
    bar.fill.solid()
    bar.fill.fore_color.rgb = palette["accent"]
    bar.line.fill.background()

    # Title & Subtitle Text Box
    tx_box = slide.shapes.add_textbox(Inches(1.4), Inches(2.0), Inches(10.8), Inches(3.5))
    tf = tx_box.text_frame
    tf.word_wrap = True
    
    p_title = tf.paragraphs[0]
    p_title.text = title
    p_title.font.name = "Segoe UI"
    p_title.font.size = Pt(44)
    p_title.font.bold = True
    p_title.font.color.rgb = palette["title_dark"]
    p_title.space_after = Pt(14)
    
    if subtitle:
        p_sub = tf.add_paragraph()
        p_sub.text = subtitle
        p_sub.font.name = "Segoe UI"
        p_sub.font.size = Pt(22)
        p_sub.font.color.rgb = palette["subtitle_dark"]
        
    if speaker_notes and hasattr(slide, "notes_slide"):
        slide.notes_slide.notes_text_frame.text = speaker_notes


def _add_slide_header(slide, title: str, palette: dict):
    """Add a clean header banner with accent bar."""
    if palette.get("is_dark_theme", False):
        slide_bg = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0), Inches(0), Inches(13.333), Inches(7.5))
        slide_bg.fill.solid()
        slide_bg.fill.fore_color.rgb = palette["dark_bg"]
        slide_bg.line.fill.background()

    header_bg = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0), Inches(0), Inches(13.333), Inches(1.2))
    header_bg.fill.solid()
    header_bg.fill.fore_color.rgb = palette["light_bg"]
    header_bg.line.fill.background()

    # Accent Bar under Header
    line = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0.8), Inches(1.15), Inches(11.733), Inches(0.04))
    line.fill.solid()
    line.fill.fore_color.rgb = palette["accent"]
    line.line.fill.background()

    # Header Title Text
    tb = slide.shapes.add_textbox(Inches(0.8), Inches(0.2), Inches(11.733), Inches(0.9))
    tf = tb.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.text = title
    p.font.name = "Segoe UI"
    p.font.size = Pt(28)
    p.font.bold = True
    p.font.color.rgb = palette["header"]


def _add_visual_recommendation(slide, visual_text: str, palette: dict, disclaimer: str = ""):
    """Add an optional visual concept card or disclaimer footer at the bottom of the slide."""
    if visual_text:
        card = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), Inches(6.3), Inches(8.5), Inches(0.8))
        card.fill.solid()
        card.fill.fore_color.rgb = palette["box_bg"]
        card.line.color.rgb = palette["accent"]

        tb = slide.shapes.add_textbox(Inches(0.9), Inches(6.35), Inches(8.3), Inches(0.7))
        tf = tb.text_frame
        tf.word_wrap = True
        p = tf.paragraphs[0]
        p.text = f"Visual Concept: {visual_text}"
        p.font.name = "Segoe UI"
        p.font.size = Pt(11)
        p.font.italic = True
        p.font.color.rgb = palette["accent"]

    if disclaimer:
        tb_disc = slide.shapes.add_textbox(Inches(9.5), Inches(6.5), Inches(3.0), Inches(0.6))
        tf_disc = tb_disc.text_frame
        tf_disc.word_wrap = True
        p_disc = tf_disc.paragraphs[0]
        p_disc.text = disclaimer
        p_disc.font.name = "Segoe UI"
        p_disc.font.size = Pt(9)
        p_disc.font.color.rgb = palette["muted"]
        p_disc.alignment = PP_ALIGN.RIGHT


def _build_bullet_slide(slide, slide_info: Dict[str, Any], palette: dict, disclaimer: str = ""):
    """Build a Standard Bullet Points Slide."""
    title = slide_info.get("title", "Key Highlights")
    _add_slide_header(slide, title, palette)
    
    tb = slide.shapes.add_textbox(Inches(0.8), Inches(1.5), Inches(11.733), Inches(4.5))
    tf = tb.text_frame
    tf.word_wrap = True
    
    content = slide_info.get("content", [])
    if isinstance(content, str):
        content = [content]

    for idx, item in enumerate(content):
        p = tf.paragraphs[0] if idx == 0 else tf.add_paragraph()
        p.text = f"•  {item}"
        p.font.name = "Segoe UI"
        p.font.size = Pt(18)
        p.font.color.rgb = palette["text_body"]
        p.space_after = Pt(14)
        
    _add_visual_recommendation(slide, slide_info.get("visual_recommendation", ""), palette, disclaimer=disclaimer)


def _build_two_column_slide(slide, slide_info: Dict[str, Any], palette: dict, disclaimer: str = ""):
    """Build a Two-Column Grid Slide."""
    title = slide_info.get("title", "Comparative Insights")
    _add_slide_header(slide, title, palette)

    col_left = slide_info.get("column_left") or slide_info.get("content", [])[:3]
    col_right = slide_info.get("column_right") or slide_info.get("content", [])[3:]

    if isinstance(col_left, str):
        col_left = [col_left]
    if isinstance(col_right, str):
        col_right = [col_right]

    # Column Left Box
    card_left = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), Inches(1.6), Inches(5.6), Inches(4.4))
    card_left.fill.solid()
    card_left.fill.fore_color.rgb = palette["box_bg"]
    card_left.line.color.rgb = palette["card_border"]

    tb_l = slide.shapes.add_textbox(Inches(1.0), Inches(1.8), Inches(5.2), Inches(4.0))
    tf_l = tb_l.text_frame
    tf_l.word_wrap = True
    for idx, item in enumerate(col_left):
        p = tf_l.paragraphs[0] if idx == 0 else tf_l.add_paragraph()
        p.text = f"•  {item}"
        p.font.name = "Segoe UI"
        p.font.size = Pt(16)
        p.font.color.rgb = palette["text_body"]
        p.space_after = Pt(12)

    # Column Right Box
    card_right = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(6.933), Inches(1.6), Inches(5.6), Inches(4.4))
    card_right.fill.solid()
    card_right.fill.fore_color.rgb = palette["box_bg"]
    card_right.line.color.rgb = palette["card_border"]

    tb_r = slide.shapes.add_textbox(Inches(7.133), Inches(1.8), Inches(5.2), Inches(4.0))
    tf_r = tb_r.text_frame
    tf_r.word_wrap = True
    for idx, item in enumerate(col_right):
        p = tf_r.paragraphs[0] if idx == 0 else tf_r.add_paragraph()
        p.text = f"•  {item}"
        p.font.name = "Segoe UI"
        p.font.size = Pt(16)
        p.font.color.rgb = palette["text_body"]
        p.space_after = Pt(12)

    _add_visual_recommendation(slide, slide_info.get("visual_recommendation", ""), palette, disclaimer=disclaimer)


def _build_quote_slide(slide, slide_info: Dict[str, Any], palette: dict, disclaimer: str = ""):
    """Build a Modern Callout / Quote Slide."""
    title = slide_info.get("title", "Key Takeaway")
    _add_slide_header(slide, title, palette)

    quote_card = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(1.5), Inches(2.0), Inches(10.333), Inches(3.6))
    quote_card.fill.solid()
    quote_card.fill.fore_color.rgb = palette["box_bg"]
    quote_card.line.color.rgb = palette["accent"]

    tb = slide.shapes.add_textbox(Inches(1.8), Inches(2.3), Inches(9.733), Inches(3.0))
    tf = tb.text_frame
    tf.word_wrap = True

    quote_text = slide_info.get("content")
    if isinstance(quote_text, list):
        quote_text = " ".join(quote_text)

    p = tf.paragraphs[0]
    p.text = f'"{quote_text}"'
    p.font.name = "Segoe UI"
    p.font.size = Pt(22)
    p.font.italic = True
    p.font.color.rgb = palette["header"]
    p.alignment = PP_ALIGN.CENTER

    _add_visual_recommendation(slide, slide_info.get("visual_recommendation", ""), palette, disclaimer=disclaimer)


def _build_metrics_slide(slide, slide_info: Dict[str, Any], palette: dict, disclaimer: str = ""):
    """Build a Multi-Metric KPI Slide."""
    title = slide_info.get("title", "Key Performance Indicators")
    _add_slide_header(slide, title, palette)

    items = slide_info.get("content", [])
    if isinstance(items, str):
        items = [items]

    card_width = Inches(3.6)
    card_height = Inches(3.8)
    gap = Inches(0.4)
    start_x = Inches(0.8)

    for i, item in enumerate(items[:3]):
        x = start_x + (card_width + gap) * i
        card = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, x, Inches(1.8), card_width, card_height)
        card.fill.solid()
        card.fill.fore_color.rgb = palette["box_bg"]
        card.line.color.rgb = palette["accent"]

        tb = slide.shapes.add_textbox(x + Inches(0.2), Inches(2.0), card_width - Inches(0.4), card_height - Inches(0.4))
        tf = tb.text_frame
        tf.word_wrap = True
        
        # Metric Number / Header
        p0 = tf.paragraphs[0]
        p0.text = f"0{i+1}"
        p0.font.name = "Segoe UI"
        p0.font.size = Pt(32)
        p0.font.bold = True
        p0.font.color.rgb = palette["accent"]
        p0.space_after = Pt(10)

        # Description
        p1 = tf.add_paragraph()
        p1.text = str(item)
        p1.font.name = "Segoe UI"
        p1.font.size = Pt(14)
        p1.font.color.rgb = palette["text_body"]

    _add_visual_recommendation(slide, slide_info.get("visual_recommendation", ""), palette, disclaimer=disclaimer)
