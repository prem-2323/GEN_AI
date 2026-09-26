"""text/document_extractor.py — text extraction for uploaded scene-source documents.

Used by ``image/routes.py`` (``/generate-scene-images-from-file``) to pull raw
text out of TXT / PDF / DOCX uploads before storyboard prompting.

All extractors accept a FastAPI ``UploadFile`` and return plain text.
"""

from __future__ import annotations

import io
import logging

from fastapi import UploadFile, HTTPException

log = logging.getLogger("gen-transform.text.document_extractor")

MAX_EXTRACT_CHARS = 200_000


def _read_bytes(upload: UploadFile) -> bytes:
    try:
        data = upload.file.read()
    finally:
        upload.file.close()
    if not data:
        raise HTTPException(status_code=400, detail="The uploaded file is empty.")
    return data


def extract_txt(upload: UploadFile) -> str:
    """Decode a plain-text upload, sniffing BOM/utf-8 with latin-1 fallback."""
    data = _read_bytes(upload)
    for encoding in ("utf-8-sig", "utf-8", "latin-1"):
        try:
            return data.decode(encoding)[:MAX_EXTRACT_CHARS]
        except UnicodeDecodeError:
            continue
    return data.decode("utf-8", errors="ignore")[:MAX_EXTRACT_CHARS]


def extract_pdf(upload: UploadFile) -> str:
    """Extract text per page from a PDF upload using pypdf."""
    try:
        from pypdf import PdfReader
    except ImportError as exc:  # pragma: no cover
        raise HTTPException(
            status_code=500, detail="PDF support requires the 'pypdf' package."
        ) from exc

    try:
        reader = PdfReader(io.BytesIO(_read_bytes(upload)))
        pages = []
        for page in reader.pages:
            text = page.extract_text() or ""
            if text.strip():
                pages.append(text)
    except HTTPException:
        raise
    except Exception as exc:
        log.warning("PDF extraction failed: %s", exc)
        raise HTTPException(
            status_code=400, detail="Could not extract text from the PDF."
        ) from exc

    return "\n\n".join(pages)[:MAX_EXTRACT_CHARS]


def extract_docx(upload: UploadFile) -> str:
    """Extract paragraph text from a DOCX upload using python-docx."""
    try:
        import docx
    except ImportError as exc:  # pragma: no cover
        raise HTTPException(
            status_code=500, detail="DOCX support requires the 'python-docx' package."
        ) from exc

    try:
        document = docx.Document(io.BytesIO(_read_bytes(upload)))
        paragraphs = [p.text for p in document.paragraphs if p.text.strip()]
        for table in document.tables:
            for row in table.rows:
                cells = [c.text.strip() for c in row.cells if c.text.strip()]
                if cells:
                    paragraphs.append(" | ".join(cells))
    except HTTPException:
        raise
    except Exception as exc:
        log.warning("DOCX extraction failed: %s", exc)
        raise HTTPException(
            status_code=400, detail="Could not extract text from the DOCX."
        ) from exc

    return "\n".join(paragraphs)[:MAX_EXTRACT_CHARS]
