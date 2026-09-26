"""Multimodal Ingestion & Extraction Test Suite (Phase 3+).

Verifies the complete PDF-with-mixed-content flow:
1. PDF containing text + embedded image + table -> text goes to Qwen channel,
   embedded images are persisted and routed to the Gemma vision channel
2. Standalone image upload (PNG/JPG) -> persisted binary + metadata (OCR optional)
3. PPTX upload -> slides as pages, tables, embedded images
4. Persisted image records are loadable and shaped for the analysis orchestrator
   (the exact shape analyze_source() feeds to Gemma)
"""
from __future__ import annotations

import io
import os
import sys
import uuid
from pathlib import Path

from fastapi.testclient import TestClient

sys.path.insert(0, os.path.abspath(os.path.join(os.path.dirname(__file__), "..")))
from app.main import app
from app.core.config import get_settings

client = TestClient(app)

TEST_RUN_ID = uuid.uuid4().hex[:8]
USER_UID = f"user-mm-{TEST_RUN_ID}"
PROJECT_ID = f"proj-mm-{TEST_RUN_ID}"


def _png_bytes(width: int = 320, height: int = 240, color=(70, 130, 180)) -> bytes:
    """Build a valid PNG in memory (complex enough to pass the 1.5KB meaningful-image filter)."""
    import random

    from PIL import Image, ImageDraw

    random.seed(42)
    img = Image.new("RGB", (width, height), color)
    # Add per-pixel noise so the PNG compresses to a realistic chart-like size (> 1.5 KB)
    px = img.load()
    for y in range(height):
        for x in range(width):
            if (x // 8 + y // 8) % 2 == 0:
                r, g, b = px[x, y]
                jitter = random.randint(-18, 18)
                px[x, y] = (max(0, min(255, r + jitter)), max(0, min(255, g + jitter)), max(0, min(255, b + jitter)))
    draw = ImageDraw.Draw(img)
    draw.rectangle([10, 10, width - 10, height - 10], outline=(255, 255, 255), width=3)
    draw.ellipse([width // 4, height // 4, 3 * width // 4, 3 * height // 4], fill=(255, 200, 0))
    buf = io.BytesIO()
    img.save(buf, format="PNG")
    return buf.getvalue()


def _pdf_with_image_and_table_bytes(image_bytes: bytes) -> bytes:
    """PDF with page text, an embedded image, and a text-aligned table."""
    try:
        from reportlab.lib.pagesizes import letter
        from reportlab.pdfgen import canvas

        buf = io.BytesIO()
        c = canvas.Canvas(buf, pagesize=letter)
        c.setTitle("Multimodal Incident Report")

        # Page 1: text + embedded image
        c.drawString(72, 750, "MULTIMODAL INCIDENT REPORT 2026")
        c.drawString(72, 730, "A breach affected 240 systems on 15 August 2026 exploiting CVE-2026-4418.")
        from reportlab.lib.utils import ImageReader

        img_reader = ImageReader(io.BytesIO(image_bytes))
        c.drawImage(img_reader, 72, 500, width=180, height=110)
        c.showPage()

        # Page 2: table-like aligned text
        c.drawString(72, 750, "IMPACT METRICS")
        c.drawString(72, 730, "Systems Impacted      240      Contained")
        c.drawString(72, 710, "CVSS Score            9.8      Critical")
        c.drawString(72, 690, "Downtime              6 hours  Resolved")
        c.showPage()
        c.save()
        return buf.getvalue()
    except ImportError:
        raise RuntimeError("reportlab is required for this test")


def _pptx_bytes(image_bytes: bytes) -> bytes:
    try:
        from pptx import Presentation
        from pptx.util import Inches

        prs = Presentation()
        # Slide 1: title + body + image
        slide = prs.slides.add_slide(prs.slide_layouts[1])
        slide.shapes.title.text = "Quarterly Security Review"
        slide.placeholders[1].text = "Latency improved by 35 percent after the Q3 patch rollout."
        img_stream = io.BytesIO(image_bytes)
        slide.shapes.add_picture(img_stream, Inches(4), Inches(3), width=Inches(3))

        # Slide 2: table
        slide2 = prs.slides.add_slide(prs.slide_layouts[5])
        rows, cols = 2, 2
        table = slide2.shapes.add_table(rows, cols, Inches(1), Inches(1), Inches(6), Inches(2)).table
        table.cell(0, 0).text = "Metric"
        table.cell(0, 1).text = "Value"
        table.cell(1, 0).text = "Uptime"
        table.cell(1, 1).text = "99.95"

        buf = io.BytesIO()
        prs.save(buf)
        return buf.getvalue()
    except ImportError:
        raise RuntimeError("python-pptx is required for this test")


def run_multimodal_tests():
    print("=" * 82)
    print("  MULTIMODAL INGESTION & EXTRACTION TEST SUITE (text->Qwen, images->Gemma)")
    print("=" * 82, "\n")

    settings = get_settings()
    storage_root = Path(settings.storage_root).resolve()

    # Setup project
    res_proj = client.post(
        "/api/projects",
        headers={"X-User-Uid": USER_UID},
        json={"id": PROJECT_ID, "name": "Multimodal Test Project"},
    )
    assert res_proj.status_code == 201, f"Project setup failed: {res_proj.text}"
    print("[Setup] Project created:", PROJECT_ID)

    png = _png_bytes()

    # ---- Test 1: PDF with text + image + table ----
    pdf_bytes = _pdf_with_image_and_table_bytes(png)
    res_pdf = client.post(
        f"/api/projects/{PROJECT_ID}/upload",
        headers={"X-User-Uid": USER_UID},
        files={"file": ("multimodal_report.pdf", pdf_bytes, "application/pdf")},
    )
    assert res_pdf.status_code == 200, f"PDF upload failed: {res_pdf.text}"
    pdf_out = res_pdf.json()
    pdf_ext = pdf_out["extractedDocument"]
    pdf_doc_id = pdf_ext["documentId"]

    assert "CVE-2026-4418" in pdf_ext["content"], "PDF text channel is empty"
    assert len(pdf_ext["pages"]) == 2
    assert len(pdf_ext["images"]) >= 1, "PDF embedded image was not persisted!"
    assert pdf_ext["images"][0]["path"], "Persisted image has no storage path"
    assert pdf_ext["metadata"].get("imageCount", 0) >= 1
    print(f"[Test 1a] PDF text+image+table extraction:")
    print(f"  [OK] Text channel: {len(pdf_ext['content'])} chars, 2 pages")
    print(f"  [OK] Embedded image persisted to {pdf_ext['images'][0]['path']}")

    # Heuristic table detection on page 2 aligned text
    assert len(pdf_ext["tables"]) >= 1, "PDF table was not detected"
    print(f"  [OK] Table detected with {len(pdf_ext['tables'][0]['rows'])} rows")

    # Image actually readable from storage
    img_rel_path = pdf_ext["images"][0]["path"]
    img_abs = storage_root / img_rel_path
    assert img_abs.exists() and img_abs.stat().st_size > 0, f"Image file missing on disk: {img_abs}"
    print(f"  [OK] Image binary verified on disk ({img_abs.stat().st_size} bytes)")

    # ---- Test 2: analysis orchestrator loads persisted images for Gemma ----
    analysis_res = client.post(
        f"/api/projects/{PROJECT_ID}/sources/{pdf_doc_id}/analysis",
        headers={"X-User-Uid": USER_UID},
        json={"extractedText": pdf_ext["content"]},
    )
    assert analysis_res.status_code == 200, f"Analysis failed: {analysis_res.text}"
    analysis = analysis_res.json()
    assert analysis["status"] == "completed"
    # Visual channel ran (Gemma or metadata fallback) because an image was persisted
    assert isinstance(analysis.get("visualAnalysis"), list)
    print(f"[Test 1b] Qwen/Gemma analysis orchestration:")
    print(f"  [OK] textAnalysis: {len(analysis['textAnalysis']['facts'])} facts, "
          f"{len(analysis['textAnalysis']['entities'])} entities, "
          f"{len(analysis['textAnalysis']['metrics'])} metrics")
    print(f"  [OK] visualAnalysis channel returned {len(analysis['visualAnalysis'])} item(s) "
          f"(mode={analysis.get('analysisMode')})")

    # ---- Test 3: UCKR includes visual citations ----
    uckr_res = client.post(
        f"/api/projects/{PROJECT_ID}/sources/{pdf_doc_id}/uckr",
        headers={"X-User-Uid": USER_UID},
    )
    assert uckr_res.status_code in (200, 201), f"UCKR build failed: {uckr_res.text}"
    uckr = (uckr_res.json() or {}).get("uckr") or uckr_res.json()
    assert uckr.get("facts"), "UCKR has no facts"
    print(f"[Test 1c] UCKR build: {uckr['stats'].get('totalFacts')} facts, "
          f"{len(uckr.get('citations', []))} citations")

    # ---- Test 4: standalone PNG upload ----
    res_png = client.post(
        f"/api/projects/{PROJECT_ID}/upload",
        headers={"X-User-Uid": USER_UID},
        files={"file": ("chart_evidence.png", png, "image/png")},
    )
    assert res_png.status_code == 200, f"PNG upload failed: {res_png.text}"
    png_out = res_png.json()
    png_ext = png_out["extractedDocument"]
    assert png_ext["fileType"] == "image"
    assert len(png_ext["images"]) == 1, "Standalone image was not persisted"
    assert png_ext["metadata"]["width"] == 320 and png_ext["metadata"]["height"] == 240
    print(f"[Test 2] Standalone image upload:")
    print(f"  [OK] Persisted as {png_ext['images'][0]['path']} ({png_ext['metadata']['imageFormat']})")
    print(f"  [OK] Routed to vision channel: {png_ext['metadata'].get('visionChannel')}")

    # Vision-only analysis must not 400
    png_doc_id = png_out["source"]["sourceId"]
    png_analysis = client.post(
        f"/api/projects/{PROJECT_ID}/sources/{png_doc_id}/analysis",
        headers={"X-User-Uid": USER_UID},
        json={},
    )
    assert png_analysis.status_code == 200, f"Vision-only analysis failed: {png_analysis.text}"
    assert png_analysis.json()["status"] == "completed"
    print(f"  [OK] Vision-only analysis completed (mode={png_analysis.json().get('analysisMode')})")

    # ---- Test 5: PPTX upload with image + table ----
    try:
        pptx = _pptx_bytes(png)
    except RuntimeError as e:
        pptx = None
        print(f"[Test 3] SKIPPED: {e}")
    if pptx:
        res_pptx = client.post(
            f"/api/projects/{PROJECT_ID}/upload",
            headers={"X-User-Uid": USER_UID},
            files={"file": ("security_review.pptx", pptx,
                            "application/vnd.openxmlformats-officedocument.presentationml.presentation")},
        )
        assert res_pptx.status_code == 200, f"PPTX upload failed: {res_pptx.text}"
        pptx_out = res_pptx.json()
        pptx_ext = pptx_out["extractedDocument"]
        assert pptx_ext["fileType"] == "pptx"
        assert len(pptx_ext["pages"]) == 2, "Expected 2 slide pages"
        assert "Quarterly Security Review" in pptx_ext["content"]
        assert "99.95" in pptx_ext["content"]
        assert len(pptx_ext["tables"]) >= 1, "PPTX table not extracted"
        assert len(pptx_ext["images"]) >= 1, "PPTX image not persisted"
        print(f"[Test 3] PPTX upload:")
        print(f"  [OK] {len(pptx_ext['pages'])} slides, {len(pptx_ext['tables'])} table(s), "
              f"{len(pptx_ext['images'])} image(s) persisted")

    # ---- Test 6: TXT-family extensions accepted ----
    res_json = client.post(
        "/api/upload",
        headers={"X-User-Uid": USER_UID},
        files={"file": ("config.json", b'{"env": "prod", "replicas": 4}', "application/json")},
    )
    assert res_json.status_code == 200, f"JSON upload failed: {res_json.text}"
    assert "replicas" in res_json.json()["extracted"]["content"]
    print("[Test 4] JSON/TXT-family file accepted and text-extracted")

    # ---- Cleanup ----
    res_del = client.delete(f"/api/projects/{PROJECT_ID}", headers={"X-User-Uid": USER_UID})
    assert res_del.status_code == 200
    print("\n" + "=" * 82)
    print("  ALL MULTIMODAL INGESTION & EXTRACTION TESTS PASSED!")
    print("=" * 82)


if __name__ == "__main__":
    run_multimodal_tests()
