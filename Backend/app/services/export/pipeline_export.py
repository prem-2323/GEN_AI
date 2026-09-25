"""Export engine (Phase 10) — deliverable content -> real files.

MongoDB stores the content; storage holds the actual files:
    storage/outputs/{uid}/{projectId}/{deliverableId}.{md,json,pptx}
The `storagePath` is recorded back on the deliverable document.
"""
from __future__ import annotations

import json
import logging
from typing import Any

from fastapi import HTTPException

from ...config.mongo import get_mongo_db
from ...utils.helpers import utcnow_iso
from ..storage.storage_service import save_output

log = logging.getLogger("gen-transform.export")


def _owner_filter(uid: str) -> dict:
    return {"$or": [{"firebaseUid": uid}, {"userId": uid}]}


def _to_markdown(dtype: str, content: dict) -> str:
    lines = [f"# {dtype.replace('_', ' ').title()}", ""]
    def _walk(obj: Any, depth: int = 2) -> None:
        if isinstance(obj, dict):
            for k, v in obj.items():
                if isinstance(v, (dict, list)):
                    lines.append(f"{'#' * min(depth, 6)} {str(k).replace('_', ' ').title()}")
                    _walk(v, depth + 1)
                else:
                    lines.append(f"- **{str(k).replace('_', ' ').title()}**: {v}")
        elif isinstance(obj, list):
            for item in obj:
                if isinstance(item, (dict, list)):
                    _walk(item, depth + 1)
                else:
                    lines.append(f"- {item}")
        else:
            lines.append(str(obj))
    _walk(content)
    return "\n".join(lines) + "\n"


def _to_pptx(dtype: str, content: dict) -> bytes:
    import io as _io
    from pptx import Presentation
    from pptx.util import Inches

    prs = Presentation()
    blank = prs.slide_layouts[6]
    if dtype == "presentation":
        for slide in (content.get("slides") or [])[:20]:
            s = prs.slides.add_slide(blank)
            title = s.shapes.add_textbox(Inches(0.5), Inches(0.3), Inches(9), Inches(1))
            title.text_frame.text = str(slide.get("title", ""))[:200]
            body = s.shapes.add_textbox(Inches(0.5), Inches(1.5), Inches(9), Inches(5))
            tf = body.text_frame
            for b in (slide.get("bullets") or [])[:8]:
                p = tf.add_paragraph()
                p.text = str(b)[:400]
                p.level = 0
    else:
        s = prs.slides.add_slide(blank)
        title = s.shapes.add_textbox(Inches(0.5), Inches(0.3), Inches(9), Inches(1))
        title.text_frame.text = dtype.replace("_", " ").title()
        body = s.shapes.add_textbox(Inches(0.5), Inches(1.5), Inches(9), Inches(5))
        tf = body.text_frame
        for line in _to_markdown(dtype, content).splitlines()[:40]:
            if line.strip():
                p = tf.add_paragraph()
                p.text = line.strip("#-* ").strip()[:300]
                p.level = 0
    buf = _io.BytesIO()
    prs.save(buf)
    return buf.getvalue()


def export_deliverable(uid: str, deliverable_id: str, fmt: str = "md") -> dict:
    from ..projects.project_service import get_project

    fmt = fmt.lower()
    if fmt not in ("md", "json", "pptx"):
        raise HTTPException(status_code=422, detail="Format must be one of: md, json, pptx.")
    db = get_mongo_db()
    doc = db["deliverables"].find_one({"deliverableId": deliverable_id, **_owner_filter(uid)})
    if not doc:
        other = db["deliverables"].find_one({"deliverableId": deliverable_id}, {"deliverableId": 1})
        if other:
            raise HTTPException(status_code=403, detail="Access denied. You do not own this deliverable.")
        raise HTTPException(status_code=404, detail="Deliverable not found.")
    doc.pop("_id", None)
    get_project(doc["projectId"], uid)

    dtype = doc.get("type", "output")
    content = doc.get("content", {})
    if fmt == "json":
        data = json.dumps({"type": dtype, "content": content}, indent=2).encode("utf-8")
    elif fmt == "pptx":
        data = _to_pptx(dtype, content if isinstance(content, dict) else {"text": content})
    else:
        data = _to_markdown(dtype, content if isinstance(content, dict) else {"text": content}).encode("utf-8")

    file_id, rel = save_output(uid, doc["projectId"], f"{deliverable_id}.{fmt}", data, deliverable_id=deliverable_id, export_type=fmt)
    
    export_id = f"exp_{deliverable_id}_{fmt}"
    mime_map = {
        "pptx": "application/vnd.openxmlformats-officedocument.presentationml.presentation",
        "docx": "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
        "pdf": "application/pdf",
        "json": "application/json",
        "md": "text/markdown",
        "mp3": "audio/mpeg",
        "mp4": "video/mp4",
    }
    export_doc = {
        "_id": export_id,
        "exportId": export_id,
        "firebaseUid": uid,
        "projectId": doc["projectId"],
        "sourceId": doc.get("sourceId"),
        "deliverableId": deliverable_id,
        "fileId": file_id,
        "type": fmt,
        "filename": f"{deliverable_id}.{fmt}",
        "mimeType": mime_map.get(fmt, "application/octet-stream"),
        "fileSize": len(data),
        "status": "completed",
        "storagePath": rel,
        "createdAt": utcnow_iso(),
    }
    db["exports"].update_one(
        {"exportId": export_id, "$or": [{"firebaseUid": uid}, {"userId": uid}]},
        {"$set": export_doc},
        upsert=True,
    )

    db["deliverables"].update_one(
        {"deliverableId": deliverable_id},
        {"$set": {"storagePath": rel, "fileId": file_id, "updatedAt": utcnow_iso(),
                  f"exports.{fmt}": rel}},
    )
    log.info("exported %s as %s -> %s (gridfs_id=%s)", deliverable_id, fmt, rel, file_id)
    return {"ok": True, "deliverableId": deliverable_id, "format": fmt,
            "storagePath": rel, "fileId": file_id, "bytes": len(data)}
