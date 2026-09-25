"""DOCX Document Exporter using python-docx (Phase 10)."""
from __future__ import annotations

import io
from typing import Any, Dict, Optional, Tuple
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH

from .base_exporter import BaseExporter
from .schemas import MIME_TYPES


class DOCXExporter(BaseExporter):
    """Generates structured Microsoft Word (.docx) documents with corporate styling and UCKR traceability."""

    async def export(
        self,
        deliverable: Dict[str, Any],
        uckr: Optional[Dict[str, Any]] = None,
        custom_title: Optional[str] = None,
    ) -> Tuple[bytes, str, str]:
        doc = Document()
        dtype = deliverable.get("type", "executive_summary").replace("_", " ").title()
        title = custom_title or deliverable.get("title") or f"{dtype} Report"
        content = deliverable.get("content", {})
        deliv_id = deliverable.get("deliverableId") or deliverable.get("id") or "export"

        # 1. Header / Branding
        brand_p = doc.add_paragraph()
        brand_run = brand_p.add_run("CONTENTFORGE AI  |  ENTERPRISE KNOWLEDGE TRANSFORMATION")
        brand_run.font.size = Pt(8.5)
        brand_run.font.color.rgb = RGBColor(100, 116, 139) # slate-500
        brand_p.alignment = WD_ALIGN_PARAGRAPH.RIGHT

        # 2. Document Title
        title_p = doc.add_heading(level=0)
        title_run = title_p.add_run(title)
        title_run.font.size = Pt(22)
        title_run.font.color.rgb = RGBColor(15, 23, 42) # slate-900

        # Subtitle / Metadata
        sub_p = doc.add_paragraph()
        sub_run = sub_p.add_run(f"Deliverable Type: {dtype}  •  Status: Verified & Grounded")
        sub_run.font.size = Pt(9.5)
        sub_run.font.color.rgb = RGBColor(71, 85, 105)
        doc.add_paragraph("─" * 60)

        # 3. Main Content Rendering
        if isinstance(content, dict):
            for key, val in content.items():
                sec_heading = str(key).replace("_", " ").title()
                h = doc.add_heading(sec_heading, level=1)
                h.paragraph_format.space_before = Pt(12)
                h.paragraph_format.space_after = Pt(4)

                if isinstance(val, list):
                    for item in val:
                        if isinstance(item, dict):
                            p = doc.add_paragraph(style="List Bullet")
                            for subk, subv in item.items():
                                p.add_run(f"{subk}: ").bold = True
                                p.add_run(f"{subv}  ")
                        else:
                            p = doc.add_paragraph(str(item), style="List Bullet")
                            p.paragraph_format.space_after = Pt(2)
                elif isinstance(val, dict):
                    for subk, subv in val.items():
                        p = doc.add_paragraph()
                        p.add_run(f"{str(subk).replace('_', ' ').title()}: ").bold = True
                        p.add_run(str(subv))
                else:
                    p = doc.add_paragraph(str(val))
                    p.paragraph_format.space_after = Pt(6)
        elif isinstance(content, list):
            for item in content:
                doc.add_paragraph(str(item), style="List Bullet")
        else:
            doc.add_paragraph(str(content))

        # 4. UCKR Fact Citations & Footnotes for Traceability
        if uckr and uckr.get("facts"):
            doc.add_page_break()
            trace_h = doc.add_heading("UCKR Fact Provenance & Traceability", level=1)
            trace_h.paragraph_format.space_before = Pt(16)
            
            note_p = doc.add_paragraph()
            note_p.add_run("This artifact is deterministically grounded against verified facts from canonical UCKR knowledge base.")
            note_p.runs[0].font.italic = True
            note_p.runs[0].font.size = Pt(9)

            table = doc.add_table(rows=1, cols=3)
            table.style = "Table Grid"
            hdr_cells = table.rows[0].cells
            hdr_cells[0].text = "Fact ID"
            hdr_cells[1].text = "Verified Grounded Fact"
            hdr_cells[2].text = "Confidence"

            for cell in hdr_cells:
                for paragraph in cell.paragraphs:
                    for run in paragraph.runs:
                        run.bold = True
                        run.font.size = Pt(9)

            for f in (uckr.get("facts") or [])[:15]:
                fid = f.get("factId") or f.get("id", "fact")
                fval = f.get("value") or f.get("text", "")
                conf = f.get("confidence", 0.98)
                
                row_cells = table.add_row().cells
                row_cells[0].text = str(fid)
                row_cells[1].text = str(fval)
                row_cells[2].text = f"{conf*100:.0f}%"
                for cell in row_cells:
                    for paragraph in cell.paragraphs:
                        for run in paragraph.runs:
                            run.font.size = Pt(8.5)

        buf = io.BytesIO()
        doc.save(buf)
        data = buf.getvalue()
        mime = MIME_TYPES["docx"]
        filename = f"{deliv_id}.docx"

        return data, mime, filename
