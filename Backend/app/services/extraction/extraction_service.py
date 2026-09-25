"""Document extraction — normalized result consumed by Phase 3+.

Normalized shape:
{
  "sourceId": "SRC...",
  "document": {"name": ..., "type": "pdf|docx|txt|md|image"},
  "text": {"content": "...", "characterCount": N},
  "pages": [{"pageNumber": N, "text": "..."}],
  "images": [{"imageId": ..., "path": ..., "pageNumber": N, "width":?, "height":?}],
  "tables": [{"tableId": ..., "pageNumber": N, "rows": [[...]], "source": "docx|heuristic"}],
  "metadata": {...}
}
"""
from __future__ import annotations

import io
import logging
import re
from typing import Any, Optional

log = logging.getLogger("gen-transform.extraction")

_WS_TABLE_RE = re.compile(r"[ \t]{3,}|\|")


def _chunk_pages(text: str, per_page: int = 2500) -> list[dict]:
    pages = []
    for i in range(0, max(1, len(text)), per_page):
        pages.append({"pageNumber": len(pages) + 1, "text": text[i : i + per_page]})
    return pages


def _heuristic_tables(pages: list[dict]) -> list[dict]:
    tables = []
    for p in pages:
        rows = [ln for ln in p["text"].splitlines() if _WS_TABLE_RE.search(ln)]
        if len(rows) >= 2:
            tables.append({
                "tableId": f"TBL-{p['pageNumber']:03d}-H1",
                "pageNumber": p["pageNumber"],
                "rows": [[c.strip() for c in re.split(r"\s{3,}|\||\t", r) if c.strip()] for r in rows[:50]],
                "source": "heuristic",
            })
    return tables


def _image_info(image_bytes: bytes) -> dict:
    try:
        from PIL import Image

        with Image.open(io.BytesIO(image_bytes)) as im:
            return {"width": im.width, "height": im.height, "format": (im.format or "").lower()}
    except Exception:
        return {}


def extract_pdf(
    file_bytes: bytes,
    filename: str,
    uid: str = "",
    project_id: str = "",
    source_id: str = "",
    persist_images: bool = False,
) -> dict:
    from pypdf import PdfReader

    reader = PdfReader(io.BytesIO(file_bytes))
    pages: list[dict] = []
    images: list[dict] = []
    for idx, page in enumerate(reader.pages):
        try:
            text = page.extract_text() or ""
        except Exception:
            text = ""
        pages.append({"pageNumber": idx + 1, "text": text.strip()})
        if persist_images:
            try:
                for img in page.images:
                    try:
                        raw = img.data
                    except Exception:
                        continue
                    from ..storage import storage_service as store

                    rel = store.save_extracted_image(uid, project_id, source_id, raw)
                    images.append({
                        "imageId": f"IMG-{idx + 1:03d}-{len(images) + 1:02d}",
                        "path": rel,
                        "pageNumber": idx + 1,
                        **_image_info(raw),
                    })
            except Exception as exc:
                log.debug("pdf image extraction skipped (p%d): %s", idx + 1, exc)
    full = "\n\n".join(p["text"] for p in pages).strip()
    
    # PyMuPDF fallback for complex or unusual PDF text encodings
    if not full:
        try:
            import pymupdf
            with pymupdf.open(stream=file_bytes, filetype="pdf") as doc:
                mu_pages = []
                for idx, page in enumerate(doc):
                    t = page.get_text("text").strip()
                    mu_pages.append({"pageNumber": idx + 1, "text": t})
                mu_full = "\n\n".join(p["text"] for p in mu_pages).strip()
                if mu_full:
                    pages = mu_pages
                    full = mu_full
        except Exception as mu_err:
            log.debug("PyMuPDF fallback skipped: %s", mu_err)

    meta = reader.metadata
    return {
        "document": {"name": filename, "type": "pdf"},
        "text": {"content": full, "characterCount": len(full)},
        "pages": pages,
        "images": images,
        "tables": _heuristic_tables(pages),
        "metadata": {
            "pageCount": len(pages),
            "title": str(getattr(meta, "title", "") or ""),
            "author": str(getattr(meta, "author", "") or ""),
        },
    }


def extract_docx(file_bytes: bytes, filename: str) -> dict:
    from docx import Document

    doc = Document(io.BytesIO(file_bytes))
    paras = [p.text for p in doc.paragraphs if p.text and p.text.strip()]
    full = "\n".join(paras).strip()
    pages = _chunk_pages(full)
    tables = []
    for ti, t in enumerate(doc.tables):
        rows = [[c.text.strip() for c in row.cells] for row in t.rows]
        if rows:
            tables.append({"tableId": f"TBL-{ti + 1:03d}", "pageNumber": 1, "rows": rows, "source": "docx"})
    if not tables:
        tables = _heuristic_tables(pages)
    core = doc.core_properties
    return {
        "document": {"name": filename, "type": "docx"},
        "text": {"content": full, "characterCount": len(full)},
        "pages": pages,
        "images": [],
        "tables": tables,
        "metadata": {
            "paragraphs": len(paras),
            "tables": len(doc.tables),
            "title": core.title or "",
            "author": core.author or "",
        },
    }


def extract_textlike(file_bytes: bytes, filename: str, kind: str) -> dict:
    text = file_bytes.decode("utf-8", errors="ignore").strip()
    pages = _chunk_pages(text)
    return {
        "document": {"name": filename, "type": kind},
        "text": {"content": text, "characterCount": len(text)},
        "pages": pages,
        "images": [],
        "tables": _heuristic_tables(pages),
        "metadata": {"lines": len(text.splitlines())},
    }


def extract_image(file_bytes: bytes, filename: str) -> dict:
    info = _image_info(file_bytes)
    return {
        "document": {"name": filename, "type": "image"},
        "text": {"content": "", "characterCount": 0},
        "pages": [],
        "images": [{"imageId": "IMG-001-01", "path": "", "pageNumber": 0, **info}],
        "tables": [],
        "metadata": {"visualOnly": True, **info},
    }


def extract_normalized(
    file_bytes: bytes,
    filename: str,
    ext: str,
    uid: str = "",
    project_id: str = "",
    source_id: str = "",
    persist_images: bool = False,
) -> dict:
    """Full normalized extraction for any supported type."""
    ext = ext.lower()
    if ext == "pdf":
        result = extract_pdf(file_bytes, filename, uid, project_id, source_id, persist_images)
    elif ext == "docx":
        result = extract_docx(file_bytes, filename)
    elif ext in ("txt", "md"):
        result = extract_textlike(file_bytes, filename, ext)
    elif ext in ("png", "jpg", "jpeg"):
        result = extract_image(file_bytes, filename)
    else:  # pragma: no cover — guarded by upload validation
        raise ValueError(f"Unsupported extension: {ext}")
    result["sourceId"] = source_id or "SRC-pending"
    return result


def extract_content(
    file_bytes: bytes, filename: str, mime_type: str = "", **kwargs: Any
) -> dict[str, Any]:
    """Legacy lightweight shape (kept for GET /api/upload backward compat)."""
    ext = filename.rsplit(".", 1)[-1].lower() if "." in filename else "txt"
    try:
        norm = extract_normalized(file_bytes, filename, ext)
        text = norm["text"]["content"]
    except Exception:
        text = file_bytes.decode("utf-8", errors="ignore").strip()
        norm = {"pages": [], "images": [], "tables": []}
    label = {"pdf": "PDF", "docx": "DOCX", "md": "TEXT"}.get(ext, ext.upper() if ext else "TEXT")
    return {
        "name": filename,
        "type": label,
        "size": f"{len(file_bytes) / 1024:.1f} KB",
        "extractedText": text,
        "status": "ready",
        "pageCount": len(norm.get("pages", [])),
        "imageCount": len(norm.get("images", [])),
        "tableCount": len(norm.get("tables", [])),
    }
