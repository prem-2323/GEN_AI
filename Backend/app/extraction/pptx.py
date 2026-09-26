"""PPTX Extraction Module (Phase 3).

Extracts from PowerPoint presentations:
- Slide-by-slide text (titles, body placeholders, grouped shapes)
- Speaker notes per slide
- Embedded tables
- Slide images (persisted for Gemma vision analysis)
- Core document metadata

Each slide becomes one "page" so downstream page-grounded citations work.
"""
from __future__ import annotations

import io
import logging
from typing import List

from ..core.exceptions import ExtractionError
from .normalizer import normalize_text
from .schemas import ExtractedDocument, ExtractedImageMeta, ExtractedPage, ExtractedSection, ExtractedTable

log = logging.getLogger("extraction.pptx")


def _shape_text(shape) -> str:
    """Recursively collect text from a slide shape (handles text frames + groups)."""
    parts: List[str] = []
    try:
        if shape.shape_type == 6:  # GROUP
            for sub in getattr(shape, "shapes", []) or []:
                parts.append(_shape_text(sub))
            return "\n".join(p for p in parts if p)
        if getattr(shape, "has_text_frame", False) and shape.text_frame is not None:
            for para in shape.text_frame.paragraphs:
                line = "".join(run.text for run in para.runs).strip()
                if line:
                    parts.append(line)
        if getattr(shape, "has_table", False) and shape.has_table:
            tbl = shape.table
            for row in tbl.rows:
                cells = [c.text.strip() for c in row.cells]
                if any(cells):
                    parts.append(" | ".join(cells))
    except Exception as exc:
        log.debug("PPTX shape text extraction warning: %s", exc)
    return "\n".join(p for p in parts if p)


def extract_pptx_document(
    file_bytes: bytes,
    filename: str,
    document_id: str,
    mime_type: str = "application/vnd.openxmlformats-officedocument.presentationml.presentation",
    uid: str = "",
    project_id: str = "",
    source_id: str = "",
    persist_images: bool = False,
) -> ExtractedDocument:
    """Extract PPTX slides, notes, tables, images, and metadata."""
    if not file_bytes:
        raise ExtractionError("PPTX file payload is empty.")

    try:
        from pptx import Presentation
    except ImportError as err:
        raise ExtractionError("python-pptx package is not installed.") from err

    try:
        prs = Presentation(io.BytesIO(file_bytes))
    except Exception as exc:
        raise ExtractionError(f"Failed to parse PPTX file '{filename}': {exc}") from exc

    pages: List[ExtractedPage] = []
    sections: List[ExtractedSection] = []
    tables: List[ExtractedTable] = []
    images: List[ExtractedImageMeta] = []
    all_text: List[str] = []

    for slide_idx, slide in enumerate(prs.slides, start=1):
        slide_lines: List[str] = []
        slide_title = ""

        for shape in slide.shapes:
            # Title first for section heading
            if shape == slide.shapes.title and shape is not None:
                try:
                    slide_title = (shape.text or "").strip()
                except Exception:
                    slide_title = ""
                continue

            # Table shapes
            try:
                if getattr(shape, "has_table", False) and shape.has_table:
                    rows_data: List[List[str]] = []
                    for row in shape.table.rows:
                        cells = [c.text.strip() for c in row.cells]
                        if any(cells):
                            rows_data.append(cells)
                    if rows_data:
                        tables.append(
                            ExtractedTable(
                                tableId=f"TBL-s{slide_idx:02d}-{len(tables) + 1:03d}",
                                pageNumber=slide_idx,
                                rows=rows_data,
                                source="pptx",
                            )
                        )
                        slide_lines.append(
                            "[Table] " + " | ".join(" / ".join(r) for r in rows_data[:3])
                        )
                    continue
            except Exception:
                pass

            # Picture shapes — persist for Gemma
            if persist_images and uid and project_id and source_id and shape.shape_type == 13:  # PICTURE
                try:
                    img = shape.image
                    img_ext = (img.ext or "png").lower()
                    if img_ext == "jpeg":
                        img_ext = "jpg"
                    from ..storage.service import save_extracted_image

                    img_id, key = save_extracted_image(
                        uid=uid,
                        project_id=project_id,
                        source_id=source_id,
                        image_bytes=img.blob,
                        ext=img_ext,
                    )
                    images.append(
                        ExtractedImageMeta(
                            imageId=img_id,
                            filename=f"{img_id}.{img_ext}",
                            path=key,
                            pageNumber=slide_idx,
                        )
                    )
                except Exception as img_err:
                    log.debug("PPTX image extraction warning (slide %d): %s", slide_idx, img_err)

            txt = _shape_text(shape)
            if txt:
                slide_lines.append(txt)

        # Speaker notes
        notes_text = ""
        try:
            if slide.has_notes_slide and slide.notes_slide is not None:
                notes_text = (slide.notes_slide.notes_text_frame.text or "").strip()
        except Exception:
            pass

        slide_body = normalize_text("\n".join(slide_lines))
        combined = slide_title + "\n" + slide_body if slide_title else slide_body
        if notes_text:
            combined += "\n[Speaker Notes] " + normalize_text(notes_text)

        if combined.strip():
            pages.append(
                ExtractedPage(
                    pageNumber=slide_idx,
                    text=combined.strip(),
                    characterCount=len(combined),
                    wordCount=len(combined.split()),
                )
            )
            all_text.append(combined.strip())
            if slide_title:
                sections.append(
                    ExtractedSection(
                        heading=slide_title,
                        level=1,
                        content=slide_body,
                    )
                )

    full_text = "\n\n".join(all_text)

    core = prs.core_properties
    metadata = {
        "title": getattr(core, "title", "") or "",
        "author": getattr(core, "author", "") or "",
        "subject": getattr(core, "subject", "") or "",
        "slideCount": len(prs.slides.__iter__.__self__._sldIdLst) if hasattr(prs.slides, "_sldIdLst") else len(pages),
        "sectionsCount": len(sections),
        "tablesCount": len(tables),
        "imageCount": len(images),
        "characterCount": len(full_text),
        "wordCount": len(full_text.split()) if full_text else 0,
    }

    return ExtractedDocument(
        documentId=document_id,
        filename=filename,
        fileType="pptx",
        mimeType=mime_type,
        sizeBytes=len(file_bytes),
        metadata=metadata,
        content=full_text,
        pages=pages,
        sections=sections,
        tables=tables,
        images=images,
        extractionStatus="completed" if (full_text or images) else "partial",
    )
