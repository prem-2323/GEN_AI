"""DOCX Extraction Module (Phase 3).

Extracts:
- Paragraphs & Headings with hierarchy
- Document sections
- Structured tables with rows & cells
- Core metadata (Title, Author, Subject, Keywords)
"""
from __future__ import annotations

import io
import logging
from typing import Any, Dict, List

from ..core.exceptions import ExtractionError
from .normalizer import normalize_text
from .schemas import (
    ExtractedDocument,
    ExtractedPage,
    ExtractedSection,
    ExtractedTable,
)

log = logging.getLogger("extraction.docx")


def extract_docx_document(
    file_bytes: bytes,
    filename: str,
    document_id: str,
    mime_type: str = "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
) -> ExtractedDocument:
    """Extract DOCX structure, paragraphs, headings, and tables."""
    if not file_bytes:
        raise ExtractionError("DOCX file payload is empty.")

    try:
        from docx import Document
    except ImportError as err:
        raise ExtractionError("python-docx package is not installed.") from err

    try:
        doc = Document(io.BytesIO(file_bytes))
    except Exception as exc:
        raise ExtractionError(f"Failed to parse DOCX file '{filename}': {exc}") from exc

    sections: List[ExtractedSection] = []
    tables: List[ExtractedTable] = []
    all_paras: List[str] = []

    current_heading = "Introduction"
    current_level = 1
    current_buffer: List[str] = []

    for p in doc.paragraphs:
        txt = (p.text or "").strip()
        if not txt:
            continue

        style_name = (getattr(p.style, "name", "") or "").lower()
        is_heading = style_name.startswith("heading") or style_name in ("title", "subtitle")

        if is_heading:
            if current_buffer:
                body = normalize_text("\n".join(current_buffer))
                if body:
                    sections.append(
                        ExtractedSection(
                            heading=current_heading,
                            level=current_level,
                            content=body,
                        )
                    )
                current_buffer = []

            current_heading = txt
            if "heading 1" in style_name or style_name == "title":
                current_level = 1
            elif "heading 2" in style_name or style_name == "subtitle":
                current_level = 2
            elif "heading 3" in style_name:
                current_level = 3
            else:
                current_level = 1
        else:
            current_buffer.append(txt)
            all_paras.append(txt)

    if current_buffer:
        body = normalize_text("\n".join(current_buffer))
        if body:
            sections.append(
                ExtractedSection(
                    heading=current_heading,
                    level=current_level,
                    content=body,
                )
            )

    # Extract tables
    for idx, table in enumerate(doc.tables):
        rows_data: List[List[str]] = []
        for row in table.rows:
            cell_texts = [c.text.strip() for c in row.cells]
            if any(cell_texts):
                rows_data.append(cell_texts)
        if rows_data:
            tables.append(
                ExtractedTable(
                    tableId=f"TBL-{idx + 1:03d}",
                    pageNumber=1,
                    rows=rows_data,
                    source="docx",
                )
            )

    full_text = normalize_text("\n\n".join(all_paras))

    # Construct estimated pages (e.g., 2000 chars per estimated page)
    pages: List[ExtractedPage] = []
    chunk_size = 2000
    for i in range(0, max(1, len(full_text)), chunk_size):
        chunk = full_text[i : i + chunk_size]
        pages.append(
            ExtractedPage(
                pageNumber=len(pages) + 1,
                text=chunk,
                characterCount=len(chunk),
                wordCount=len(chunk.split()) if chunk else 0,
            )
        )

    # Core metadata
    core = doc.core_properties
    metadata = {
        "title": getattr(core, "title", "") or "",
        "author": getattr(core, "author", "") or "",
        "subject": getattr(core, "subject", "") or "",
        "keywords": getattr(core, "keywords", "") or "",
        "paragraphsCount": len(all_paras),
        "sectionsCount": len(sections),
        "tablesCount": len(tables),
        "characterCount": len(full_text),
        "wordCount": len(full_text.split()) if full_text else 0,
    }

    return ExtractedDocument(
        documentId=document_id,
        filename=filename,
        fileType="docx",
        mimeType=mime_type,
        sizeBytes=len(file_bytes),
        metadata=metadata,
        content=full_text,
        pages=pages,
        sections=sections,
        tables=tables,
        images=[],
        extractionStatus="completed",
    )
